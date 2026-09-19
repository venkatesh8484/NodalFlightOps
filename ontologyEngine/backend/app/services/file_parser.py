import csv
import io
import json
import xml.etree.ElementTree as ET

import fitz
from openpyxl import load_workbook
from docx import Document


def _decode_bytes(raw: bytes) -> str:
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("latin-1", errors="ignore")


def extract_text(filename: str, raw: bytes) -> tuple[str, str | None]:
    ext = filename.lower().rsplit(".", 1)
    ext = f".{ext[1]}" if len(ext) == 2 else ""

    if ext in {".txt", ".md", ".doc", ".docs"}:
        return _decode_bytes(raw), None

    if ext == ".json":
        parsed = json.loads(_decode_bytes(raw))
        return json.dumps(parsed, indent=2), None

    if ext == ".xml":
        root = ET.fromstring(raw)
        text_chunks: list[str] = []
        for elem in root.iter():
            if elem.text and elem.text.strip():
                text_chunks.append(elem.text.strip())
        return "\n".join(text_chunks), None

    if ext == ".csv":
        text = _decode_bytes(raw)
        reader = csv.reader(io.StringIO(text))
        lines = [" | ".join(row) for row in reader]
        return "\n".join(lines), None

    if ext == ".xlsx":
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            parts: list[str] = []
            for sheet in wb.worksheets:
                parts.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(v) for v in row if v is not None and str(v).strip()]
                    if values:
                        parts.append(" | ".join(values))
            return "\n".join(parts), None
        except Exception:
            return _decode_bytes(raw), "XLSX parser fallback used"

    if ext == ".docx":
        try:
            document = Document(io.BytesIO(raw))
            lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            return "\n".join(lines), None
        except Exception:
            return _decode_bytes(raw), "DOCX parser fallback used"

    if ext == ".pdf":
        doc = fitz.open(stream=raw, filetype="pdf")
        text = " ".join(page.get_text("text") for page in doc)
        return text.strip(), None

    return "", "Unsupported extension for parser"
