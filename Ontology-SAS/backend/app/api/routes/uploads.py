from pathlib import Path
import os
from typing import Any

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook

from app.core.config import ALLOWED_EXTENSIONS, STATIC_JSON_DIR, STATIC_XLSX_DIR, UPLOAD_DIR
from app.schemas.upload import (
    OntoGPTOutput,
    ProcessedFile,
    StageOutputs,
    TaxonomySummary,
    UploadResponse,
)
from app.services.data_extractor import extract_from_file, merge_extractions
# from app.services.ontogpt_service import RealOntoGPTPipeline
from app.services.stage_xlsx_export import export_stage_excels
from app.services.SixPhaseOntologyPipeline import SixPhaseOntologyPipeline

router = APIRouter(prefix="/api/v1/uploads", tags=["uploads"])


def _limit_json_payload(
    value: Any,
    *,
    max_array_items: int,
    max_object_keys: int,
    max_depth: int,
    depth: int = 0,
) -> tuple[Any, bool]:
    if depth >= max_depth:
        return "...truncated...", True

    if isinstance(value, dict):
        items = list(value.items())
        truncated = len(items) > max_object_keys
        limited: dict[str, Any] = {}
        for key, item in items[:max_object_keys]:
            limited_value, child_truncated = _limit_json_payload(
                item,
                max_array_items=max_array_items,
                max_object_keys=max_object_keys,
                max_depth=max_depth,
                depth=depth + 1,
            )
            limited[key] = limited_value
            truncated = truncated or child_truncated
        return limited, truncated

    if isinstance(value, list):
        truncated = len(value) > max_array_items
        limited_list = []
        for item in value[:max_array_items]:
            limited_value, child_truncated = _limit_json_payload(
                item,
                max_array_items=max_array_items,
                max_object_keys=max_object_keys,
                max_depth=max_depth,
                depth=depth + 1,
            )
            limited_list.append(limited_value)
            truncated = truncated or child_truncated
        return limited_list, truncated

    return value, False


def _resolve_artifact_path(raw_path: str) -> Path:
    path = Path(raw_path)
    if not path.is_absolute():
        path = Path.cwd() / path
    resolved = path.resolve()

    allowed_roots = [UPLOAD_DIR.resolve(), (UPLOAD_DIR.parent / "generated_outputs").resolve()]
    if not any(str(resolved).startswith(str(root)) for root in allowed_roots):
        raise HTTPException(status_code=400, detail="Artifact path is not allowed.")
    if not resolved.exists():
        raise HTTPException(status_code=404, detail="Artifact not found.")
    return resolved


def _build_latest_stage_outputs() -> StageOutputs | None:
    json_paths = {
        "controlled_vocabulary_json": STATIC_JSON_DIR / "01_controlled_vocabulary.json",
        "metadata_standard_json": STATIC_JSON_DIR / "02_metadata_standard.json",
        "taxonomy_json": STATIC_JSON_DIR / "03_taxonomy.json",
        "thesaurus_json": STATIC_JSON_DIR / "04_thesaurus.json",
        "ontology_json": STATIC_JSON_DIR / "05_ontology.json",
        "knowledge_graph_jsonld": STATIC_JSON_DIR / "06_knowledge_graph.jsonld",
    }
    xlsx_paths = {
        "controlled_vocabulary_xlsx": STATIC_XLSX_DIR / "01_controlled_vocabulary.xlsx",
        "metadata_xlsx": STATIC_XLSX_DIR / "02_metadata.xlsx",
        "taxonomy_xlsx": STATIC_XLSX_DIR / "03_taxonomy.xlsx",
        "thesaurus_xlsx": STATIC_XLSX_DIR / "04_thesaurus.xlsx",
        "ontology_xlsx": STATIC_XLSX_DIR / "05_ontology.xlsx",
        "knowledge_graph_xlsx": STATIC_XLSX_DIR / "06_knowledge_graph.xlsx",
    }

    if not any(path.exists() for path in [*json_paths.values(), *xlsx_paths.values()]):
        return None

    payload = {
        "run_folder": str(STATIC_JSON_DIR.parent),
        **{key: str(path) for key, path in xlsx_paths.items()},
        **{key: (str(path) if path.exists() else None) for key, path in json_paths.items()},
    }
    return StageOutputs(**payload)


def _build_latest_taxonomy_summary(stage_outputs: StageOutputs | None) -> TaxonomySummary | None:
    if not stage_outputs:
        return None

    taxonomy_path = Path(stage_outputs.taxonomy_json) if stage_outputs.taxonomy_json else None
    kg_path = Path(stage_outputs.knowledge_graph_jsonld) if stage_outputs.knowledge_graph_jsonld else None

    concept_count = 0
    if taxonomy_path and taxonomy_path.exists():
        import json

        with open(taxonomy_path, "r", encoding="utf-8") as handle:
            taxonomy_data = json.load(handle)
            if isinstance(taxonomy_data, list):
                concept_count = len(taxonomy_data)

    return TaxonomySummary(
        concept_count=concept_count,
        output_json=str(taxonomy_path) if taxonomy_path else "",
        output_skos_jsonld=str(kg_path) if kg_path else "",
    )


@router.get("/latest", response_model=UploadResponse)
def latest_outputs():
    stage_outputs = _build_latest_stage_outputs()
    if not stage_outputs:
        return UploadResponse(
            pipeline_mode="current",
            accepted=0,
            rejected=0,
            files=[],
            taxonomy=None,
            stage_outputs=None,
            ontogpt_output=None,
        )

    return UploadResponse(
        pipeline_mode="current",
        accepted=0,
        rejected=0,
        files=[],
        taxonomy=_build_latest_taxonomy_summary(stage_outputs),
        stage_outputs=stage_outputs,
        ontogpt_output=None,
    )


@router.get("/artifact-preview")
def preview_artifact(
    path: str,
    max_graph_nodes: int = 80,
    max_array_items: int = 300,
    max_object_keys: int = 120,
    max_depth: int = 8,
):
    artifact_path = _resolve_artifact_path(path)
    suffix = artifact_path.suffix.lower()

    if suffix in {".json", ".jsonld"}:
        import json

        with open(artifact_path, "r", encoding="utf-8") as handle:
            data = json.load(handle)

        truncated = False
        total_graph_nodes = None
        shown_graph_nodes = None

        if isinstance(data, dict) and isinstance(data.get("@graph"), list):
            graph = data["@graph"]
            total_graph_nodes = len(graph)
            if len(graph) > max_graph_nodes:
                data = dict(data)
                data["@graph"] = graph[:max_graph_nodes]
                shown_graph_nodes = max_graph_nodes
                truncated = True
            else:
                shown_graph_nodes = len(graph)

        limited_data, was_limited = _limit_json_payload(
            data,
            max_array_items=max_array_items,
            max_object_keys=max_object_keys,
            max_depth=max_depth,
        )
        truncated = truncated or was_limited

        return {
            "path": str(artifact_path),
            "kind": "json",
            "data": limited_data,
            "truncated": truncated,
            "total_graph_nodes": total_graph_nodes,
            "shown_graph_nodes": shown_graph_nodes,
        }

    if suffix in {".txt", ".md", ".csv", ".xml", ".yaml", ".yml"}:
        return {
            "path": str(artifact_path),
            "kind": "text",
            "data": artifact_path.read_text(encoding="utf-8", errors="ignore"),
        }

    if suffix == ".xlsx":
        workbook = load_workbook(artifact_path, data_only=True)
        sheets = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(["" if value is None else value for value in row])
                if len(rows) >= 40:
                    break
            sheets.append(
                {
                    "name": sheet.title,
                    "rows": rows,
                }
            )
        return {
            "path": str(artifact_path),
            "kind": "xlsx",
            "sheets": sheets,
        }

    raise HTTPException(status_code=400, detail="Unsupported artifact type.")


@router.post("/process", response_model=UploadResponse)
async def process_uploads(
    files: list[UploadFile] = File(...),
    output_base: str = Form(default="company_taxonomy"),
    pipeline_mode: str = Form(default="current"),
):
    accepted_files: list[ProcessedFile] = []
    rejected_count = 0
    merged_text_parts: list[str] = []
    stage_outputs = None
    ontogpt_output = None

    for upload in files:
        suffix = Path(upload.filename).suffix.lower()
        content = await upload.read()

        if suffix not in ALLOWED_EXTENSIONS:
            rejected_count += 1
            accepted_files.append(
                ProcessedFile(
                    filename=upload.filename,
                    extension=suffix,
                    bytes=len(content),
                    parsed=False,
                    text_length=0,
                    warning="Unsupported extension",
                )
            )
            continue

        target_path = UPLOAD_DIR / upload.filename
        target_path.write_bytes(content)

        try:
            text, warning = extract_from_file(upload.filename, content)
            if text:
                merged_text_parts.append((upload.filename, text))
                accepted_files.append(
                    ProcessedFile(
                        filename=upload.filename,
                        extension=suffix,
                        bytes=len(content),
                        parsed=True,
                        text_length=len(text),
                        warning=warning,
                    )
                )
            else:
                accepted_files.append(
                    ProcessedFile(
                        filename=upload.filename,
                        extension=suffix,
                        bytes=len(content),
                        parsed=False,
                        text_length=0,
                        warning=warning or "No text extracted",
                    )
                )
        except Exception as exc:
            accepted_files.append(
                ProcessedFile(
                    filename=upload.filename,
                    extension=suffix,
                    bytes=len(content),
                    parsed=False,
                    text_length=0,
                    warning=f"Parse failed: {exc}",
                )
            )

    taxonomy_summary = None

    # Merge all file extractions: raw text (phases 5-6) + noun-only (phases 1-4)
    merged_raw, merged_nouns = merge_extractions(merged_text_parts)

    # Save both versions for inspection
    output_dir = UPLOAD_DIR.parent / "generated_outputs"
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "merged_input_nouns.txt").write_text(merged_nouns, encoding="utf-8")
    (output_dir / "merged_input_raw.txt").write_text(merged_raw, encoding="utf-8")

    # Check that at least one LLM provider is configured
    from app.services.llm_helper import MODEL_PROVIDER
    has_databricks = os.getenv("DATABRICKS_HOST") and os.getenv("DATABRICKS_TOKEN")
    has_azure = os.getenv("AZURE_OPENAI_KEY") and os.getenv("AZURE_OPENAI_ENDPOINT") and os.getenv("AZURE_OPENAI_DEPLOYMENT_ID")
    has_openai = os.getenv("OPENAI_API_KEY")

    if not has_databricks and not has_azure and not has_openai:
        raise HTTPException(
            status_code=400,
            detail=(
                "AI mode requires at least one provider configured:\n"
                "  • Databricks: DATABRICKS_HOST + DATABRICKS_TOKEN\n"
                "  • Azure OpenAI: AZURE_OPENAI_KEY + AZURE_OPENAI_ENDPOINT + AZURE_OPENAI_DEPLOYMENT_ID\n"
                "  • OpenAI: OPENAI_API_KEY\n"
                f"Current MODEL_PROVIDER={MODEL_PROVIDER}"
            ),
        )

    if not merged_nouns.strip():
        raise HTTPException(
            status_code=400,
            detail="No parseable text found in uploaded files. AI pipeline cannot run.",
        )

    if pipeline_mode == "ontogpt":
        raise HTTPException(
            status_code=400,
            detail="OntoGPT mode is temporarily disabled. Use pipeline_mode=current.",
        )

        # OntoGPT mode is intentionally disabled for now.
        # verify_ssl = os.getenv("DATABRICKS_VERIFY_SSL", "false").lower() == "true"
        # service = RealOntoGPTPipeline(
        #     databricks_token=resolved_token,
        #     databricks_host=resolved_host,
        #     model_name=resolved_model,
        #     verify_ssl=verify_ssl,
        # )
        # try:
        #     ontogpt_result = service.run_pipeline(
        #         raw_text=successful_text,
        #         output_base=output_base,
        #     )
        # except Exception as exc:
        #     raise HTTPException(
        #         status_code=503,
        #         detail=f"AI pipeline failed: {exc}",
        #     ) from exc
        # ontogpt_output = OntoGPTOutput(**ontogpt_result["stage_json_outputs"])

    if merged_nouns.strip():
        service = SixPhaseOntologyPipeline()
        try:
            stage_result = service.run_pipeline(
                noun_text=merged_nouns,
                raw_text=merged_raw,
                metadata_rows=[file.model_dump() for file in accepted_files],
                output_base=output_base,
            )
        except Exception as exc:
            import traceback
            raise HTTPException(
                status_code=503,
                detail=f"AI pipeline failed: {type(exc).__name__}: {exc}\n{traceback.format_exc()}",
            ) from exc
        taxonomy_summary = TaxonomySummary(
            concept_count=stage_result["concept_count"],
            output_json=stage_result["output_json"],
            output_skos_jsonld=stage_result["output_skos_jsonld"],
        )

        xlsx_paths = export_stage_excels(
            output_base=output_base,
            controlled_vocab=stage_result["controlled_vocab"],
            metadata_standard=stage_result["metadata_standard"],
            taxonomy=stage_result["taxonomy"],
            thesaurus=stage_result["thesaurus"],
            ontology=stage_result["ontology"],
            knowledge_graph=stage_result["knowledge_graph"],
        )
        xlsx_paths.update(
            {
                "controlled_vocabulary_json": stage_result["stage_json_outputs"]["controlled_vocabulary_json"],
                "metadata_standard_json": stage_result["stage_json_outputs"]["metadata_standard_json"],
                "taxonomy_json": stage_result["stage_json_outputs"]["taxonomy_json"],
                "thesaurus_json": stage_result["stage_json_outputs"]["thesaurus_json"],
                "ontology_json": stage_result["stage_json_outputs"]["ontology_json"],
                "knowledge_graph_jsonld": stage_result["stage_json_outputs"]["knowledge_graph_jsonld"],
            }
        )
        stage_outputs = StageOutputs(**xlsx_paths)

    return UploadResponse(
        pipeline_mode=pipeline_mode,
        accepted=sum(1 for f in accepted_files if f.parsed),
        rejected=rejected_count,
        files=accepted_files,
        taxonomy=taxonomy_summary,
        stage_outputs=stage_outputs,
        ontogpt_output=ontogpt_output,
    )
