from __future__ import annotations

import re
from pathlib import Path
import json

from openpyxl import Workbook

from app.core.config import STATIC_XLSX_DIR

# openpyxl rejects control characters (except tab/newline/CR)
_ILLEGAL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]") 

def _sanitize(value: str) -> str:
    return _ILLEGAL_CHARS_RE.sub("", value)


def _safe_sheet_row(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return _sanitize(", ".join(str(item) for item in value))
    return _sanitize(str(value))


def export_stage_excels(
    output_base: str,
    controlled_vocab: list[dict],
    metadata_standard: list[dict],
    taxonomy: list[dict],
    thesaurus: list[dict],
    ontology: dict,
    knowledge_graph: dict,
) -> dict:
    run_dir = STATIC_XLSX_DIR
    run_dir.mkdir(parents=True, exist_ok=True)

    controlled_vocab_path = run_dir / "01_controlled_vocabulary.xlsx"
    metadata_path = run_dir / "02_metadata.xlsx"
    taxonomy_path = run_dir / "03_taxonomy.xlsx"
    thesaurus_path = run_dir / "04_thesaurus.xlsx"
    ontology_path = run_dir / "05_ontology.xlsx"
    knowledge_graph_path = run_dir / "06_knowledge_graph.xlsx"

    for path in [
        controlled_vocab_path,
        metadata_path,
        taxonomy_path,
        thesaurus_path,
        ontology_path,
        knowledge_graph_path,
    ]:
        if path.exists():
            path.unlink()

    _write_controlled_vocabulary(controlled_vocab_path, controlled_vocab)
    _write_metadata(metadata_path, metadata_standard)
    _write_taxonomy(taxonomy_path, taxonomy)
    _write_thesaurus(thesaurus_path, thesaurus)
    _write_ontology(ontology_path, ontology)
    _write_knowledge_graph(knowledge_graph_path, knowledge_graph)

    return {
        "run_folder": str(run_dir),
        "controlled_vocabulary_xlsx": str(controlled_vocab_path),
        "metadata_xlsx": str(metadata_path),
        "taxonomy_xlsx": str(taxonomy_path),
        "thesaurus_xlsx": str(thesaurus_path),
        "ontology_xlsx": str(ontology_path),
        "knowledge_graph_xlsx": str(knowledge_graph_path),
    }


def _write_controlled_vocabulary(path: Path, concepts: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Controlled Vocabulary"
    ws.append(["Approved_Term", "Aliases"])

    for concept in concepts:
        ws.append(
            [
                _sanitize(concept.get("Approved_Term", "")),
                _safe_sheet_row(concept.get("Aliases", [])),
            ]
        )

    wb.save(path)


def _write_metadata(path: Path, metadata_rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Metadata Standard"
    ws.append([
        "Concept_ID",
        "Preferred_Term_PT",
        "Used_For_UF",
        "Status",
    ])

    for row in metadata_rows:
        ws.append(
            [
                _sanitize(row.get("Concept_ID", "")),
                _sanitize(row.get("Preferred_Term_PT", "")),
                _safe_sheet_row(row.get("Used_For_UF", [])),
                _sanitize(row.get("Status", "")),
            ]
        )

    wb.save(path)


def _write_taxonomy(path: Path, concepts: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Taxonomy"
    ws.append(
        [
            "Concept_ID",
            "Preferred_Term_PT",
            "Broader_Term_BT",
            "Narrower_Term_NT",
            "Related_Term_RT",
        ]
    )

    for concept in concepts:
        ws.append(
            [
                _sanitize(concept.get("Concept_ID", "")),
                _sanitize(concept.get("Preferred_Term_PT", "")),
                _safe_sheet_row(concept.get("Broader_Term_BT", [])),
                _safe_sheet_row(concept.get("Narrower_Term_NT", [])),
                _safe_sheet_row(concept.get("Related_Term_RT", [])),
            ]
        )

    wb.save(path)


def _write_thesaurus(path: Path, rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Thesaurus"
    ws.append([
        "Concept_ID",
        "PT",
        "UF",
        "BT",
        "NT",
        "Related_Term_RT",
        "Scope_Note_SN",
    ])

    for row in rows:
        ws.append(
            [
                _sanitize(row.get("Concept_ID", "")),
                _sanitize(row.get("PT", "")),
                _safe_sheet_row(row.get("UF", [])),
                _safe_sheet_row(row.get("BT", [])),
                _safe_sheet_row(row.get("NT", [])),
                _safe_sheet_row(row.get("Related_Term_RT", [])),
                _sanitize(row.get("Scope_Note_SN", "")),
            ]
        )

    wb.save(path)


def _write_ontology(path: Path, ontology: dict) -> None:
    wb = Workbook()
    ws_classes = wb.active
    ws_classes.title = "Classes"
    ws_classes.append(["Class"])

    for class_name in ontology.get("Classes", []):
        ws_classes.append([str(class_name)])

    ws_props = wb.create_sheet("Object_Properties")
    ws_props.append(["Verb", "Domain", "Range"])
    for prop in ontology.get("Object_Properties", []):
        if not isinstance(prop, dict):
            continue
        ws_props.append([
            prop.get("Verb", ""),
            prop.get("Domain", ""),
            prop.get("Range", ""),
        ])

    wb.save(path)


def _write_knowledge_graph(path: Path, graph_doc: dict) -> None:
    wb = Workbook()
    ws_context = wb.active
    ws_context.title = "Context"
    ws_context.append(["key", "id", "type"])

    context = graph_doc.get("@context", {})
    if isinstance(context, dict):
        for key, value in context.items():
            if isinstance(value, dict):
                ws_context.append([key, value.get("@id", ""), value.get("@type", "")])
            else:
                ws_context.append([key, str(value), ""])

    ws_graph = wb.create_sheet("Graph")
    ws_graph.append(["@id", "@type", "properties"])
    for node in graph_doc.get("@graph", []):
        if not isinstance(node, dict):
            continue
        node_id = node.get("@id", "")
        node_type = node.get("@type", "")
        props = {k: v for k, v in node.items() if k not in {"@id", "@type"}}
        ws_graph.append([node_id, node_type, json.dumps(props, ensure_ascii=False)])

    wb.save(path)
