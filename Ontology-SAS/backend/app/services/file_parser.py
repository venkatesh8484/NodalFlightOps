import csv
import io
import json
import re
import xml.etree.ElementTree as ET

import fitz
from openpyxl import load_workbook
from docx import Document


MAX_CATEGORY_VALUES = 20
MAX_SAMPLED_ROWS = 1200


def _is_number_like(value: str) -> bool:
    return bool(re.fullmatch(r"[-+]?\d+(?:[.,]\d+)?", value.strip()))


def _is_date_like(value: str) -> bool:
    v = value.strip()
    return bool(
        re.fullmatch(r"\d{4}[-/]\d{1,2}[-/]\d{1,2}", v)
        or re.fullmatch(r"\d{1,2}[-/]\d{1,2}[-/]\d{2,4}", v)
    )


def _is_long_free_text(value: str) -> bool:
    return len(value.strip()) > 80


def _collect_categorical_values(values: list[str]) -> list[str]:
    cleaned: list[str] = []
    seen: set[str] = set()
    for raw in values:
        value = str(raw).strip()
        if not value:
            continue
        key = value.lower()
        if key in seen:
            continue
        seen.add(key)
        cleaned.append(value)

    if len(cleaned) < 2 or len(cleaned) > MAX_CATEGORY_VALUES:
        return []

    mostly_numeric_or_dates = sum(1 for v in cleaned if _is_number_like(v) or _is_date_like(v)) >= max(2, len(cleaned) - 1)
    if mostly_numeric_or_dates:
        return []

    if sum(1 for v in cleaned if _is_long_free_text(v)) > 1:
        return []

    return cleaned


def _summarize_tabular_dataset(dataset_name: str, headers: list[str], rows: list[list[str]]) -> str:
    valid_headers = [str(h).strip() for h in headers if str(h).strip()]
    if not valid_headers:
        return f"Dataset: {dataset_name}\nNo usable column headers found."

    columns: dict[int, list[str]] = {idx: [] for idx in range(len(valid_headers))}
    sampled = 0
    for row in rows:
        if sampled >= MAX_SAMPLED_ROWS:
            break
        sampled += 1
        for idx in range(len(valid_headers)):
            value = row[idx] if idx < len(row) else ""
            columns[idx].append(str(value).strip())

    lines = [f"Dataset: {dataset_name}", "Use schema and categorical values only.", "Column Headers:"]
    for header in valid_headers:
        lines.append(f"- {header}")

    lines.append("Categorical Reference Values:")
    any_categories = False
    for idx, header in enumerate(valid_headers):
        enums = _collect_categorical_values(columns.get(idx, []))
        if enums:
            any_categories = True
            lines.append(f"- {header}: {', '.join(enums[:MAX_CATEGORY_VALUES])}")

    if not any_categories:
        lines.append("- none detected")

    return "\n".join(lines)


def _parse_sql_schema(sql_text: str) -> str:
    create_blocks = re.findall(r"CREATE\s+TABLE\s+.*?;", sql_text, flags=re.IGNORECASE | re.DOTALL)
    alter_fk_blocks = re.findall(
        r"ALTER\s+TABLE\s+.*?FOREIGN\s+KEY\s*\(.*?\)\s*REFERENCES\s+.*?;",
        sql_text,
        flags=re.IGNORECASE | re.DOTALL,
    )

    lines = [
        "SQL Schema Summary:",
        "Use table names, column names, and foreign-key relations only.",
    ]

    table_count = 0
    for block in create_blocks:
        table_match = re.search(
            r"CREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?[`\"\[]?([\w.]+)[`\"\]]?",
            block,
            flags=re.IGNORECASE,
        )
        if not table_match:
            continue

        table_count += 1
        table_name = table_match.group(1)
        lines.append(f"Table: {table_name}")

        body_match = re.search(r"\((.*)\)", block, flags=re.DOTALL)
        if not body_match:
            continue

        body = body_match.group(1)
        for raw_line in body.splitlines():
            line = raw_line.strip().rstrip(",")
            if not line:
                continue
            upper = line.upper()

            if "FOREIGN KEY" in upper and "REFERENCES" in upper:
                fk_match = re.search(
                    r"FOREIGN\s+KEY\s*\(([^)]+)\)\s*REFERENCES\s+[`\"\[]?([\w.]+)[`\"\]]?\s*\(([^)]+)\)",
                    line,
                    flags=re.IGNORECASE,
                )
                if fk_match:
                    src_cols = fk_match.group(1).strip()
                    ref_table = fk_match.group(2).strip()
                    ref_cols = fk_match.group(3).strip()
                    lines.append(f"- FK: ({src_cols}) -> {ref_table}({ref_cols})")
                continue

            if upper.startswith(("PRIMARY KEY", "UNIQUE", "INDEX", "KEY", "CONSTRAINT", "CHECK")):
                continue

            col_match = re.match(r"[`\"\[]?([A-Za-z_][\w$]*)[`\"\]]?\s+(.+)", line)
            if not col_match:
                continue

            col_name = col_match.group(1)
            col_type = col_match.group(2).split()[0]
            lines.append(f"- Column: {col_name} ({col_type})")

    for block in alter_fk_blocks:
        alter_match = re.search(
            r"ALTER\s+TABLE\s+[`\"\[]?([\w.]+)[`\"\]]?.*?FOREIGN\s+KEY\s*\(([^)]+)\)\s*REFERENCES\s+[`\"\[]?([\w.]+)[`\"\]]?\s*\(([^)]+)\)",
            block,
            flags=re.IGNORECASE | re.DOTALL,
        )
        if not alter_match:
            continue
        src_table = alter_match.group(1).strip()
        src_cols = alter_match.group(2).strip()
        ref_table = alter_match.group(3).strip()
        ref_cols = alter_match.group(4).strip()
        lines.append(f"FK (ALTER): {src_table}({src_cols}) -> {ref_table}({ref_cols})")

    if table_count == 0:
        return "SQL Schema Summary:\nNo CREATE TABLE statements found; provide DDL for best ontology quality."

    return "\n".join(lines)


def _decode_bytes(raw: bytes) -> str:
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("latin-1", errors="ignore")


def extract_text(filename: str, raw: bytes) -> tuple[str, str | None]:
    ext = filename.lower().rsplit(".", 1)
    ext = f".{ext[1]}" if len(ext) == 2 else ""

    if ext in {".txt", ".md", ".doc", ".docs"}:
        return _decode_bytes(raw), None

    if ext == ".sql":
        sql_text = _decode_bytes(raw)
        return _parse_sql_schema(sql_text), None

    if ext == ".json":
        parsed = json.loads(_decode_bytes(raw))
        return json.dumps(parsed, indent=2), None

    if ext == ".xml":
        root = ET.fromstring(raw)
        text_chunks: list[str] = []
        for elem in root.iter():
            if elem.text and elem.text.strip():
                text_chunks.append(elem.text.strip())
        return "\n".join(text_chunks), None

    if ext == ".csv":
        text = _decode_bytes(raw)
        reader = csv.reader(io.StringIO(text))
        rows = [row for row in reader if any(str(cell).strip() for cell in row)]
        if not rows:
            return "Dataset: CSV\nNo rows found.", None
        headers = rows[0]
        data_rows = rows[1:]
        return _summarize_tabular_dataset("CSV", headers, data_rows), None

    if ext == ".xlsx":
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            parts: list[str] = []
            for sheet in wb.worksheets:
                rows: list[list[str]] = []
                for row in sheet.iter_rows(values_only=True):
                    rows.append(["" if v is None else str(v) for v in row])
                    if len(rows) >= MAX_SAMPLED_ROWS + 1:
                        break

                if not rows:
                    continue

                headers = rows[0]
                data_rows = rows[1:]
                parts.append(_summarize_tabular_dataset(f"XLSX::{sheet.title}", headers, data_rows))

            if not parts:
                return "Workbook contains no readable sheets.", "XLSX parser fallback used"
            return "\n".join(parts), None
        except Exception:
            return _decode_bytes(raw), "XLSX parser fallback used"

    if ext == ".docx":
        try:
            document = Document(io.BytesIO(raw))
            lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            return "\n".join(lines), None
        except Exception:
            return _decode_bytes(raw), "DOCX parser fallback used"

    if ext == ".pdf":
        doc = fitz.open(stream=raw, filetype="pdf")
        text = " ".join(page.get_text("text") for page in doc)
        return text.strip(), None

    return "", "Unsupported extension for parser"
