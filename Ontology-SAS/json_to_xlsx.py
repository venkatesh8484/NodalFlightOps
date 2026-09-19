"""
Standalone JSON / JSON-LD -> XLSX converter.

Usage:
    python json_to_xlsx.py <input.json|input.jsonld> [output.xlsx]
    python json_to_xlsx.py <input_folder> [output_folder]

If the input is a single file and no output path is given, the .xlsx file is
written next to the input file with the same name (extension replaced with .xlsx).

If the input is a folder, every .json/.jsonld file directly inside it is
converted to its own .xlsx file. If no output folder is given, files are
written to an "ontooutput" folder created next to this script.

Handles:
- A plain list of flat dicts          -> one sheet, columns = union of all keys
- A dict of named lists (e.g. Phase 5 ontology: {"Classes": [...], "Object_Properties": [...]})
                                       -> one sheet per top-level list key
- JSON-LD with an "@graph" array      -> a "Graph" sheet from @graph, plus a
                                          "Context" sheet from "@context" if present
- A single flat dict                  -> one sheet with Key / Value columns

Nested lists/dicts inside a cell are flattened to a readable string
(list of primitives -> comma-joined, dict / list of dicts -> compact JSON)
so nothing gets silently dropped.

Requires: openpyxl (already installed in backend/requirements.txt).
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook

_ILLEGAL_CHARS_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_SHEET_NAME_ILLEGAL_RE = re.compile(r"[\\/*?:\[\]]")


def _sanitize(value: str) -> str:
    return _ILLEGAL_CHARS_RE.sub("", value)


def _cell_value(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str):
        return _sanitize(value)
    if isinstance(value, (list, tuple)):
        if all(isinstance(v, (str, int, float, bool)) or v is None for v in value):
            return _sanitize(", ".join("" if v is None else str(v) for v in value))
        return _sanitize(json.dumps(value, ensure_ascii=False))
    if isinstance(value, dict):
        return _sanitize(json.dumps(value, ensure_ascii=False))
    return _sanitize(str(value))


def _safe_sheet_title(name: str) -> str:
    name = _SHEET_NAME_ILLEGAL_RE.sub("_", str(name))[:31]
    return name or "Sheet"


def _write_list_sheet(wb: Workbook, title: str, rows: list, first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = _safe_sheet_title(title)

    dict_rows = [r for r in rows if isinstance(r, dict)]
    if not dict_rows:
        ws.append(["Value"])
        for row in rows:
            ws.append([_cell_value(row)])
        return ws

    columns: list[str] = []
    for row in dict_rows:
        for key in row.keys():
            if key not in columns:
                columns.append(key)

    ws.append(columns)
    for row in dict_rows:
        ws.append([_cell_value(row.get(col)) for col in columns])
    return ws


def _write_dict_sheet(wb: Workbook, title: str, data: dict, first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = _safe_sheet_title(title)
    ws.append(["Key", "Value"])
    for key, value in data.items():
        ws.append([_sanitize(str(key)), _cell_value(value)])
    return ws


def convert(input_path: Path, output_path: Path) -> Path:
    data = json.loads(input_path.read_text(encoding="utf-8"))

    wb = Workbook()
    wrote_first = False

    if isinstance(data, list):
        _write_list_sheet(wb, input_path.stem or "Data", data, first=True)
        wrote_first = True

    elif isinstance(data, dict):
        if "@graph" in data and isinstance(data["@graph"], list):
            _write_list_sheet(wb, "Graph", data["@graph"], first=True)
            wrote_first = True
            if isinstance(data.get("@context"), dict):
                _write_dict_sheet(wb, "Context", data["@context"])
            for key, value in data.items():
                if key in ("@graph", "@context"):
                    continue
                if isinstance(value, list):
                    _write_list_sheet(wb, key, value)
                elif isinstance(value, dict):
                    _write_dict_sheet(wb, key, value)
        else:
            list_keys = [k for k, v in data.items() if isinstance(v, list)]
            if list_keys:
                for key in list_keys:
                    _write_list_sheet(wb, key, data[key], first=not wrote_first)
                    wrote_first = True
                remaining = {k: v for k, v in data.items() if k not in list_keys}
                if remaining:
                    _write_dict_sheet(wb, "Info", remaining)
            else:
                _write_dict_sheet(wb, input_path.stem or "Data", data, first=True)
                wrote_first = True
    else:
        _write_dict_sheet(wb, "Data", {"Value": data}, first=True)
        wrote_first = True

    if not wrote_first:
        wb.active.title = "Sheet1"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def convert_folder(input_dir: Path, output_dir: Path) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    written = []
    for file_path in sorted(input_dir.iterdir()):
        if file_path.is_file() and file_path.suffix.lower() in (".json", ".jsonld"):
            output_path = output_dir / (file_path.stem + ".xlsx")
            written.append(convert(file_path, output_path))
    return written


def main() -> None:
    if len(sys.argv) < 2:
        print("Usage: python json_to_xlsx.py <input.json|input.jsonld|input_folder> [output.xlsx|output_folder]")
        sys.exit(1)

    input_path = Path(sys.argv[1]).resolve()
    if not input_path.exists():
        print(f"Input path not found: {input_path}")
        sys.exit(1)

    if input_path.is_dir():
        output_dir = Path(sys.argv[2]).resolve() if len(sys.argv) >= 3 else Path(__file__).resolve().parent / "ontooutput"
        results = convert_folder(input_path, output_dir)
        if not results:
            print(f"No .json/.jsonld files found in: {input_path}")
            sys.exit(1)
        for result in results:
            print(f"Wrote: {result}")
        return

    output_path = Path(sys.argv[2]).resolve() if len(sys.argv) >= 3 else input_path.with_suffix(".xlsx")
    result = convert(input_path, output_path)
    print(f"Wrote: {result}")


if __name__ == "__main__":
    main()
